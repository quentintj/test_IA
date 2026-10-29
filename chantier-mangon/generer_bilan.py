"""Génère le bilan électricité Mangon (Word) : budgété (devis), signé (plans), constaté (29/10/2026)."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from donnees import CONSTATS, LECTURE, PIECES

ICI = Path(__file__).resolve().parent
PHOTOS = ICI / "photos"
SORTIE = ICI / "Bilan_electricite_Mangon_2026-10-29.docx"

STATUTS = {"ok": "Conforme", "nok": "Non conforme", "abs": "Absent", "": "À vérifier", None: "Non vérifié"}
COUL_STATUT = {"ok": "D9EAD3", "nok": "F9DDC5", "abs": "F4CCCC", "": "FFF2CC", None: "EEEEEE"}
CATEGORIES = {
    "M": ("Manquements de l'électricien", "Prévu au devis et/ou aux plans signés, mais absent ou non conforme sur place.", "F4CCCC"),
    "C": ("Points à clarifier", "Le devis, les plans et le constat ne concordent pas : il faut la réponse de l'électricien.", "FFF2CC"),
    "S": ("Écarts au budget (suppléments ou moins-values)", "Posé ou demandé hors devis, ou posé en quantité différente du devis.", "DCE6F2"),
}


def ombrer(cellule, couleur):
    tc_pr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc_pr.append(shd)


def ecrire(cellule, texte, gras=False, taille=8.5, italique=False):
    cellule.text = ""
    run = cellule.paragraphs[0].add_run(texte)
    run.bold, run.italic = gras, italique
    run.font.size = Pt(taille)


def tableau(doc, entetes, largeurs):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, titre in enumerate(entetes):
        ecrire(t.rows[0].cells[i], titre, gras=True)
        ombrer(t.rows[0].cells[i], "D9D9D9")
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    entete = OxmlElement("w:tblHeader")
    entete.set(qn("w:val"), "true")
    tr_pr.append(entete)
    for i, l in enumerate(largeurs):
        t.columns[i].width = Cm(l)
    return t, largeurs


def ligne(t, largeurs, valeurs):
    cells = t.add_row().cells
    for i, l in enumerate(largeurs):
        cells[i].width = Cm(l)
    for i, v in enumerate(valeurs):
        if v is not None:
            ecrire(cells[i], v)
    return cells


def points():
    """Tous les points de la checklist, enrichis du constat et de la lecture."""
    out = []
    for pid, piece, alt, items in PIECES:
        for iid, libelle, devis, plans, ecart in items:
            cle = f"{pid}--{iid}"
            statut, remarque, photos = CONSTATS.get(cle, (None, "", []))
            cat, lecture = LECTURE.get(cle, (None, ""))
            out.append(dict(cle=cle, piece=piece, alt=alt, libelle=libelle, devis=devis, plans=plans,
                            ecart=bool(ecart), statut=statut, remarque=remarque, photos=photos, cat=cat, lecture=lecture))
    return out


def main():
    pts = points()
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    for marge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, marge, Cm(1.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("Chantier Mangon – Vérification électricité", level=0)
    p = doc.add_paragraph()
    p.add_run("Visite du 29/10/2026").bold = True
    p.add_run(" · 86 points de contrôle · 14 photos")
    doc.add_paragraph(
        "Chaque point est repris depuis le début : ce qui a été budgété (devis de l'électricien), "
        "ce qui a été signé (plans électriques signés) et ce qui a été constaté sur place. "
        "La colonne « Lecture » en déduit s'il s'agit d'un manquement, d'un écart au budget ou d'un point à clarifier.")

    # 1. Synthèse
    doc.add_heading("1. Synthèse", level=1)
    t, lg = tableau(doc, ["Constat du 29/10", "Nombre"], [6, 2.5])
    for code in ["abs", "nok", "", "ok", None]:
        n = sum(1 for x in pts if x["statut"] == code)
        c = ligne(t, lg, [STATUTS[code], str(n)])
        ombrer(c[0], COUL_STATUT[code])
    ligne(t, lg, ["Total", str(len(pts))])
    doc.add_paragraph()
    t, lg = tableau(doc, ["Lecture budget / signé / constaté", "Nombre", "Points"], [6, 2.5, 16.5])
    for code, (titre, _, coul) in CATEGORIES.items():
        sel = [x for x in pts if x["cat"] == code]
        c = ligne(t, lg, [titre, str(len(sel)), " · ".join(f"{x['libelle']} ({x['piece']})" for x in sel)])
        ombrer(c[0], coul)

    doc.add_heading("Attention : numérotation des chambres", level=2)
    doc.add_paragraph(
        "Le devis et les plans ne numérotent pas les chambres de la même façon (le listing du tableau "
        "diffère aussi du devis). Dans ce document, la numérotation suit le devis : chambre 1 = 14,12 m² "
        "(« Ch 4 » sur les plans), chambre 2 = 16,13 m² (« Ch 3 »), chambre 3 = 16,63 m² (« Ch 2 »), "
        "chambre 4 = 11,69 m² (« Ch 1 », absente du devis). La commande de la lampe depuis les têtes de lit "
        "et les prises USB sont prévues au devis dans la chambre 3, et non dans la chambre 1. C'est "
        "probablement l'origine des remarques « c'était dans la chambre 1 ».")

    # 2-4. Une section par catégorie
    entetes = ["Pièce", "Élément", "Budgété (devis)", "Signé (plans)", "Constaté 29/10", "Lecture / action", "Photo"]
    largeurs = [2.8, 3.2, 3.2, 3.6, 3.6, 6.4, 3.2]
    for n, (code, (titre, expl, coul)) in enumerate(CATEGORIES.items(), start=2):
        doc.add_heading(f"{n}. {titre}", level=1)
        doc.add_paragraph(expl)
        t, lg = tableau(doc, entetes, largeurs)
        for x in (x for x in pts if x["cat"] == code):
            constat = STATUTS[x["statut"]] + (f" — « {x['remarque']} »" if x["remarque"] else "")
            c = ligne(t, lg, [x["piece"], x["libelle"], x["devis"], x["plans"], constat, x["lecture"], None])
            ecrire(c[0], x["piece"], gras=True)
            ombrer(c[4], COUL_STATUT[x["statut"]])
            if x["photos"]:
                par = c[6].paragraphs[0]
                for ph in x["photos"]:
                    par.add_run().add_picture(str(PHOTOS / f"{ph}.jpg"), width=Cm(2.6))
            else:
                ecrire(c[6], "–")

    # 5. Écarts relevés sur les plans mais constatés conformes
    n = len(CATEGORIES) + 2
    doc.add_heading(f"{n}. Écarts devis / plans constatés conformes sur place", level=1)
    doc.add_paragraph("Ces points différaient entre le devis et les plans signés, mais ont été cochés conformes le 29/10. "
                      "À garder en tête pour le décompte final.")
    t, lg = tableau(doc, ["Pièce", "Élément", "Budgété (devis)", "Signé (plans)"], [5, 6, 7, 8])
    for x in pts:
        if x["ecart"] and x["statut"] == "ok" and not x["cat"]:
            ligne(t, lg, [x["piece"], x["libelle"], x["devis"], x["plans"]])

    # 6. Non vérifiés
    n += 1
    doc.add_heading(f"{n}. Points non vérifiés le 29/10", level=1)
    t, lg = tableau(doc, ["Pièce", "Élément", "Budgété (devis)", "Signé (plans)"], [5, 6, 7, 8])
    for x in pts:
        if x["statut"] is None:
            ligne(t, lg, [x["piece"], x["libelle"], x["devis"], x["plans"]])

    # 7. Conformes
    n += 1
    doc.add_heading(f"{n}. Points conformes", level=1)
    par_piece = {}
    for x in pts:
        if x["statut"] == "ok" and not x["cat"]:
            par_piece.setdefault(x["piece"], []).append(x["libelle"])
    for piece, libs in par_piece.items():
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(piece + " : ").bold = True
        p.add_run(", ".join(libs))

    # 8. Photos
    n += 1
    doc.add_section().orientation = WD_ORIENT.PORTRAIT
    s2 = doc.sections[-1]
    s2.page_width, s2.page_height = s.page_height, s.page_width
    doc.add_heading(f"{n}. Photos de la visite du 29/10/2026", level=1)
    num = 0
    for x in pts:
        for ph in x["photos"]:
            num += 1
            doc.add_paragraph().add_run().add_picture(str(PHOTOS / f"{ph}.jpg"), height=Cm(10.5))
            cap = doc.add_paragraph()
            cap.add_run(f"Photo {num} – {x['piece']} – {x['libelle']}").bold = True
            cap.add_run(f" · constat : {STATUTS[x['statut']].lower()}" + (f" — « {x['remarque']} »" if x["remarque"] else ""))

    doc.save(SORTIE)
    print(SORTIE)


if __name__ == "__main__":
    main()
