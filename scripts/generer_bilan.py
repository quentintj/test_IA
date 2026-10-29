"""Génère le bilan de vérification électricité du chantier Mangon (Word)."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SORTIE = Path(__file__).resolve().parent.parent / "chantier-mangon" / "Bilan_electricite_Mangon_2026-10-29.docx"

A_COMPLETER = "à compléter"

# (pièce, élément, statut constaté, remarque, nb photos)
POINTS = [
    ("Tableau électrique", "Différentiels", "Conforme", "Problème de disjoncteur ?", 1),
    ("Hall d'entrée et hall cuisine", "Spots au plafond", "Conforme", "Attention prolongement", 2),
    ("Salle de douche", "Point au-dessus de la douche", "À vérifier",
     "2 prises électriques. Un interrupteur du mauvais côté ? Les prises s'activent si la lumière est allumée", 1),
    ("Garde-manger", "Spots au plafond", "À vérifier", "Uniquement 2", 1),
    ("Cuisine", "Spots au plafond", "À vérifier", "6 spots", 1),
    ("Cuisine", "Interrupteur des spots", "Non conforme", "Mauvais endroit", 1),
    ("Cuisine", "Point lumineux sous meuble", "Absent", "", 0),
    ("Cuisine", "Prises doubles plan de travail", "Conforme", "Il y avait également une prise prévue ici", 1),
    ("Buanderie", "Appliques extérieures", "Non conforme", "", 1),
    ("Buanderie", "Prise extérieure", "Non conforme", "", 1),
    ("Salon", "Appliques murales, 2 directions", "Absent", "", 0),
    ("Salle à manger", "Interrupteurs près de la porte du hall", "Conforme", "Attention erreur", 1),
    ("Garage", "Tube LED 1,20 m", "Non conforme", "Au plus haut", 0),
    ("Garage", "Prise NIKO Hydro noire", "Non conforme", "Posée en gris", 0),
    ("Garage", "Alimentation chaudière", "Non conforme", "", 0),
    ("Garage", "Alimentations pompe à chaleur (int. + ext.) et chauffe-eau", "Non conforme", "", 0),
    ("Garage", "Éclairage extérieur côté jardin", "Absent", "De quoi parle-t-on ?", 0),
    ("Escalier et hall des chambres 1 et 2", "Éclairage au plafond", "Non conforme",
     "5 spots, comme on a toujours discuté", 1),
    ("Hall des chambres 3 et 4", "Point sur rail", "Non conforme", "Quid des rails ?", 1),
    ("Hall des chambres 3 et 4", "Applique, 2 directions + témoin", "Non conforme",
     "Quid des témoins ? Quid des deux directions ? Quid de l'applique ?", 0),
    ("Salle de bains", "Liaison équipotentielle", "À vérifier", "?", 0),
    ("Chambre 1", "Prises doubles", "Conforme",
     "Gros attention : les appliques et les interrupteurs à côté du lit doivent pouvoir éteindre la lampe de la chambre", 0),
    ("Chambre 2", "Prises doubles", "À vérifier", "Que 2 prises doubles", 0),
    ("Chambre 3", "Prises simples", "Absent", "", 0),
    ("Chambre 3", "Prises USB A / USB C", "À vérifier", "C'était dans la chambre 1", 0),
    ("Chambre 4", "Prises doubles", "À vérifier", "Il y en a 3", 0),
    ("Bureau / dressing", "Point lumineux + prise", "Absent", "", 0),
    ("Combles", "Prises doubles", "À vérifier", "?", 0),
]

COULEURS = {
    "Absent": "F4CCCC",
    "Non conforme": "F9DDC5",
    "À vérifier": "FFF2CC",
    "Conforme": "D9EAD3",
}
ORDRE_STATUTS = ["Absent", "Non conforme", "À vérifier", "Conforme"]


def ombrer(cellule, couleur):
    tc_pr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc_pr.append(shd)


def repeter_entete(ligne):
    tr_pr = ligne._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def ecrire(cellule, texte, gras=False, taille=8.5, italique=False, couleur=None):
    cellule.text = ""
    run = cellule.paragraphs[0].add_run(texte)
    run.bold = gras
    run.italic = italique
    run.font.size = Pt(taille)
    if couleur:
        run.font.color.rgb = couleur


def tableau(doc, entetes, largeurs):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, titre in enumerate(entetes):
        c = t.rows[0].cells[i]
        ecrire(c, titre, gras=True, taille=8.5)
        ombrer(c, "D9D9D9")
    repeter_entete(t.rows[0])
    t.autofit = False
    for i, l in enumerate(largeurs):
        t.columns[i].width = Cm(l)
    for ligne in t.rows:
        for i, l in enumerate(largeurs):
            ligne.cells[i].width = Cm(l)
    return t


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for marge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, marge, Cm(1.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("Chantier Mangon – Vérification électricité", level=0)
    p = doc.add_paragraph()
    p.add_run("Date de la visite : ").bold = True
    p.add_run("29/10/2026")
    p = doc.add_paragraph()
    p.add_run("Objet : ").bold = True
    p.add_run("repartir du début — ce qui a été budgété, ce qui a été signé, ce qui a été constaté, "
              "et les manquements qui en découlent.")

    note = doc.add_paragraph()
    r = note.add_run("Document provisoire : les colonnes « Budgété » et « Signé (devis) » sont à compléter. "
                     "Le budget électricité, le devis signé et les photos du 29/10 n'étaient pas encore disponibles "
                     "au moment de la rédaction.")
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # Synthèse
    doc.add_heading("1. Synthèse", level=1)
    comptes = {s: [pt for pt in POINTS if pt[2] == s] for s in ORDRE_STATUTS}
    t = tableau(doc, ["Statut", "Nombre", "Points concernés"], [3.5, 2, 19.5])
    for statut in ORDRE_STATUTS:
        cells = t.add_row().cells
        ecrire(cells[0], statut, gras=True, taille=9)
        ombrer(cells[0], COULEURS[statut])
        ecrire(cells[1], str(len(comptes[statut])), taille=9)
        ecrire(cells[2], " · ".join(f"{pt[1]} ({pt[0]})" for pt in comptes[statut]), taille=8.5)
    total = t.add_row().cells
    ecrire(total[0], "Total", gras=True, taille=9)
    ecrire(total[1], str(len(POINTS)), gras=True, taille=9)
    ecrire(total[2], f"{sum(pt[4] for pt in POINTS)} photos prises lors de la visite", taille=8.5)

    # Tableau comparatif
    doc.add_heading("2. Tableau comparatif par pièce", level=1)
    doc.add_paragraph("Écart à déterminer pour chaque point : manquement de l'électricien (prévu au devis signé), "
                      "supplément (non prévu au devis) ou point à clarifier.")
    entetes = ["Pièce", "Élément", "Budgété", "Signé (devis)", "Constaté 29/10", "Remarque", "Photos", "Écart / action"]
    largeurs = [3.2, 4.2, 2.2, 2.2, 2.4, 6.2, 1.4, 3.2]
    t = tableau(doc, entetes, largeurs)
    piece_prec = None
    for piece, element, statut, remarque, photos in POINTS:
        cells = t.add_row().cells
        ecrire(cells[0], piece if piece != piece_prec else "", gras=True)
        piece_prec = piece
        ecrire(cells[1], element)
        for i in (2, 3, 7):
            ecrire(cells[i], A_COMPLETER, italique=True, couleur=RGBColor(0x80, 0x80, 0x80))
        ecrire(cells[4], statut, gras=True)
        ombrer(cells[4], COULEURS[statut])
        ecrire(cells[5], remarque)
        ecrire(cells[6], str(photos) if photos else "–")
        for i, l in enumerate(largeurs):
            cells[i].width = Cm(l)

    # Questions ouvertes
    doc.add_heading("3. Questions à poser à l'électricien", level=1)
    for q in [
        "Tableau : quel est le problème de disjoncteur constaté ?",
        "Halls : comment le prolongement des spots a-t-il été réalisé ?",
        "Salle de douche : pourquoi les prises sont-elles commandées par l'interrupteur de la lumière ? "
        "L'interrupteur est-il du bon côté ?",
        "Garde-manger : 2 spots posés — combien étaient prévus ?",
        "Cuisine : 6 spots posés — combien étaient prévus ? Où est la prise supplémentaire prévue au plan de travail ?",
        "Garage : de quel éclairage extérieur côté jardin parle-t-on ?",
        "Hall des chambres 3 et 4 : où sont les rails, les témoins, la commande 2 directions et l'applique ?",
        "Salle de bains : la liaison équipotentielle a-t-elle été réalisée et contrôlée ?",
        "Chambre 1 : les appliques et interrupteurs près du lit commandent-ils la lampe de la chambre ?",
        "Chambres 2 et 4 : nombre de prises doubles prévu (2 posées en chambre 2, 3 en chambre 4) ?",
        "Chambre 3 : prises USB A / USB C — chambre 1 ou chambre 3 ?",
        "Combles : prises doubles prévues ?",
    ]:
        doc.add_paragraph(q, style="List Number")

    # Documents manquants
    doc.add_heading("4. Documents à joindre pour finaliser", level=1)
    for d in [
        "Budget électricité initial (estimation de départ)",
        "Devis signé de l'électricien et éventuels avenants",
        "Plan électrique convenu",
        f"Photos de la visite du 29/10/2026 ({sum(pt[4] for pt in POINTS)} photos référencées dans le tableau)",
    ]:
        doc.add_paragraph(d, style="List Bullet")

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SORTIE)
    print(SORTIE)


if __name__ == "__main__":
    main()
